"""
Generate Conclusions section as a formatted Word document.
Writes to: ../docs/Conclusions_Trends_2026.docx
Target: ~400 words.
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm

OUT = Path(__file__).parent.parent / "docs" / "Conclusions_Trends_2026.docx"

doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(3)


def heading(text, level=1):
    p = doc.add_heading(text, level=level)
    p.runs[0].font.size = Pt(12 if level == 1 else 11)
    return p


def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.runs[0]
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    return p


# ============================================================
# CONCLUSIONS SECTION
# ============================================================

heading("6. Conclusions", level=1)

body(
    "This review set out to map the landscape of drivers, trends, and signals "
    "shaping the future of fisheries, drawing on a three-step search strategy "
    "that combined a structured review of peer-reviewed fisheries futures "
    "literature, a cross-sectoral scan of megatrend foresight publications, "
    "and an initial horizon scan for weak and emerging signals. Taken together, "
    "the 206 sources compiled across the three strategies constitute a "
    "structured evidence base for ongoing monitoring of the forces shaping "
    "fisheries systems over the medium to long term."
)

body(
    "The central empirical finding of this review is a pronounced disciplinary "
    "gap in the coverage of technological drivers within fisheries futures research. "
    "Whilst technological change — encompassing automation, digitalisation, "
    "artificial intelligence, and synthetic biology — accounts for 28% of entries "
    "in the cross-sectoral megatrend literature, no comparable body of futures-oriented "
    "work was identified within the manually screened fisheries-specific literature "
    "(0%; gap = +28 percentage points). Social and economic drivers are similarly "
    "underrepresented relative to the cross-sectoral baseline, whilst environmental "
    "drivers — particularly climate change and ecological disruption — dominate "
    "the fisheries futures evidence base (53% of classified entries). These "
    "findings suggest that fisheries research has yet to integrate fully the "
    "breadth of transformative forces that broader foresight scholarship has "
    "identified as structurally significant."
)

body(
    "The review makes three practical contributions. First, it provides a "
    "structured, STEEP-classified database of 206 sources that can serve as a "
    "reference point and working document for researchers, educators, and "
    "policymakers engaged with the future of fisheries. Second, it identifies "
    "specific thematic and methodological gaps that should inform research "
    "priorities: the absence of technological futures research is the most "
    "acute of these, but the relative neglect of social and economic foresight "
    "methodologies in fisheries science also merits sustained attention. Third, "
    "it introduces a protocol for ongoing weak signal monitoring that, if "
    "maintained and expanded collaboratively, could provide an early warning "
    "function for emerging developments not yet captured in the formal literature."
)

body(
    "Future work should prioritise completion of full-text screening of the "
    "720 candidate papers identified via the automated OpenAlex search, which "
    "will substantially expand the evidence base and may revise the gap analysis "
    "reported here. Expanding the signal database, developing inter-rater "
    "reliability protocols for the STEEP classification, and extending the "
    "search to non-English-language foresight literature — particularly from "
    "East and Southeast Asian research communities — represent further priorities. "
    "Ultimately, the findings presented here argue for a more deliberate and "
    "structured engagement between fisheries science and the broader futures "
    "studies field: one that draws on the full range of foresight methods and "
    "the full spectrum of megatrends, rather than the environmental framing "
    "that currently predominates."
)

# ── Save ──────────────────────────────────────────────────────
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"Saved: {OUT}")
