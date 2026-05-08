"""
Generate Introduction section as a formatted Word document.
Writes to: ../docs/Introduction_Trends_2026.docx
Target: ~800 words.

Notes:
- Fixes Hitunen -> Hiltunen (misspelling in original draft)
- Operationalises three research questions (addresses reviewer major concern #7)
- Uses only citations present in original draft
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm

OUT = Path(__file__).parent.parent / "docs" / "Introduction_Trends_2026.docx"

doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(3)


def heading(text, level=1):
    p = doc.add_heading(text, level=level)
    p.runs[0].font.size = Pt(12 if level == 1 else 11)
    return p


def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.runs[0]
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    return p


def italic_note(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.name = "Times New Roman"
    p.paragraph_format.space_after = Pt(4)
    return p


# ============================================================
# INTRODUCTION SECTION
# ============================================================

heading("1. Introduction", level=1)

body(
    "Fisheries serve as a vital pillar for food security, nutrition, and employment "
    "within societies worldwide, providing livelihoods for an estimated 600 million "
    "people and constituting a primary source of animal protein for over three billion "
    "(FAO, 2022). Yet fisheries systems are subject to an accelerating array of "
    "structural pressures — from climate-driven changes in ocean productivity and "
    "species distribution, to shifting patterns of global demand, technological "
    "disruption, and evolving governance frameworks — that collectively challenge "
    "the sustainability and adaptability of the sector. The capacity to anticipate "
    "these pressures, rather than merely respond to them, depends in part on the "
    "availability of structured, forward-looking analyses of the forces most likely "
    "to shape fisheries over the medium to long term. Yet systematic horizon scanning "
    "and trend monitoring have, until recently, remained peripheral to mainstream "
    "fisheries research."
)

body(
    "The futures studies field offers a repertoire of conceptual tools and methods "
    "for the structured analysis of change. Drivers of change are forces or factors "
    "that shift conditions within a system, often operating at scales beyond the "
    "immediate control of individual stakeholders, but with direct and sometimes "
    "rapid consequences for how institutions, communities, and businesses function "
    "(Vecchiato and Roveda, 2010). Trends are widespread, directional changes that "
    "unfold over time, influencing attitudes, policies, and economic activity at "
    "societal scale; they emerge as specific innovations or cultural shifts become "
    "broadly recognised and institutionalised (Celi and Colombi, 2020). When trends "
    "operate at a global scale and over multi-decadal timescales, they are commonly "
    "described as megatrends — powerful forces expected to bring about significant "
    "transformations in societal or natural environments over a period of five to "
    "twenty or more years (Naisbitt, 1982; Hajkowicz, 2015). Weak signals, by "
    "contrast, are early and often ambiguous indicators of potentially significant "
    "future changes: subtle, socially embedded observations that, when interpreted "
    "through an appropriate analytical lens, can alert practitioners to developments "
    "that do not yet register in formal datasets or policy agendas "
    "(Hiltunen, 2010; Rossel, 2021)."
)

body(
    "Research on megatrends has grown substantially over the past two decades, "
    "with contributions from business consultancies (PWC, 2016; KPMG, 2014; EY, 2020), "
    "institutional research organisations (Knowledge4Policy, 2022; Naughtin et al., "
    "2022, 2024; Linthorst and de Waal, 2020), and academic journals across multiple "
    "disciplines (Deml et al., 2022; Kalaitzi et al., 2020; Malik and Janowska, 2021). "
    "These analyses identify a convergent set of transformative forces — including "
    "climate change, demographic shifts, technological acceleration, and economic "
    "power redistribution — that are expected to reshape societies and industries "
    "across sectors. A comparable analytical tradition is beginning to emerge within "
    "fisheries research. Makino et al. (2011) identified five key drivers structuring "
    "the future of Japanese fisheries, while a growing body of scenario-based and "
    "trend-analytical work has examined the implications of specific environmental "
    "and governance changes for fisheries systems globally (Garcia and Grainger, 2005; "
    "Garcia and Rosenberg, 2010; Pauly et al., 2003; Sumaila et al., 2016; "
    "Schickele et al., 2020). More recently, narrative scenario methods have been "
    "applied in marine socio-ecological system (MSES) research to explore potential "
    "futures under combined social and ecological pressures (Merrie et al., 2018; "
    "Planque et al., 2019; Spijkers et al., 2018, 2021; Lübker et al., 2023; "
    "Levin et al., 2013)."
)

body(
    "Despite these advances, futures studies methodologies remain comparatively "
    "underutilised in fisheries research relative to other domains such as urban "
    "development, energy, and public health. In particular, the systematic integration "
    "of cross-sectoral foresight knowledge — including the full range of megatrends "
    "identified across technology, economy, and social domains — into fisheries "
    "futures analysis has not yet been attempted. This represents both a methodological "
    "gap and a practical risk: governance and management frameworks developed without "
    "reference to the broader landscape of transformative change may be poorly "
    "equipped to navigate the structural shifts that fisheries systems are likely "
    "to encounter over the coming decades."
)

body(
    "This paper addresses this gap through a three-step horizon scanning and "
    "structured literature review. The study is guided by three research questions: "
    "(1) What drivers and trends are documented in the peer-reviewed futures literature "
    "specific to fisheries, and how are they distributed across the Social, "
    "Technological, Economic, Environmental, and Political (STEEP) framework? "
    "(2) How does the thematic coverage of the fisheries-specific futures literature "
    "compare to that of the cross-sectoral megatrend foresight evidence base, and "
    "what disciplinary gaps can be identified? "
    "(3) What emerging signals in the broader futures landscape may be relevant to "
    "fisheries systems but are not yet represented in the fisheries futures literature? "
    "To address these questions, the study compiled a structured database of 206 "
    "sources — combining peer-reviewed fisheries futures literature, cross-sectoral "
    "megatrend publications, and an initial horizon scan for weak and emerging signals "
    "— and introduces this database as a living document intended for ongoing "
    "collaborative expansion by the research community."
)

body(
    "The paper is structured as follows. Section 2 introduces the conceptual "
    "framework and key terminology. Section 3 describes the three-step search "
    "strategy and classification procedure. Section 4 presents the results organised "
    "by drivers, trends, and signals. Section 5 discusses the implications of the "
    "findings for researchers, educators, and policymakers, and identifies "
    "priorities for future work. Section 6 concludes."
)

# ── Save ──────────────────────────────────────────────────────
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"Saved: {OUT}")
