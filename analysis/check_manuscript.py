"""Quick word count and structure check for the merged manuscript."""
from pathlib import Path
from docx import Document

doc = Document(Path(__file__).parent.parent / "docs" / "Manuscript_Trends_2026.docx")

words = sum(len(p.text.split()) for p in doc.paragraphs)
chars = sum(len(p.text) for p in doc.paragraphs)
print(f"Word count : {words:,}")
print(f"Char count : {chars:,}")
print(f"Paragraphs : {len(doc.paragraphs)}")
print()
print("=== STRUCTURE ===")
for p in doc.paragraphs:
    txt = p.text.strip()
    if not txt:
        continue
    style = p.style.name
    if style.startswith("Heading"):
        level = style.replace("Heading ", "")
        indent = "  " * (int(level) - 1) if level.isdigit() else ""
        print(f"{indent}[H{level}] {txt[:90]}")
    elif any(r.bold for r in p.runs if r.text.strip()) and len(txt) < 80:
        # Bold short lines = likely section labels or title
        print(f"  [bold] {txt[:80]}")
