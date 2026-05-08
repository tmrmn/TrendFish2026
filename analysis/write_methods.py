"""
Generate Methods section as a formatted Word document.
Writes to: ../docs/Methods_Trends_2026.docx
Target: ~800 words.

Replaces the placeholder bullet-point structure in 'Method Section Trends.docx'
with complete prose. All search parameters, dates, databases, inclusion/exclusion
criteria, and the STEEP classification procedure are documented.

PRISMA 2020 flow numbers:
  Strategy 1 (automated, OpenAlex): 2,100 identified -> 479 duplicates removed
    -> 1,621 screened -> 901 excluded -> 720 pending full-text
  Strategy 1 (manual): 60 included
  Strategy 2 (grey lit): 145 included
  Strategy 3 (signals): ongoing living document

Key citations (all present in original draft):
  Page et al. 2021          — PRISMA 2020
  Bearman et al. 2012       — SLR
  Pati & Lorusso 2017       — SLR
  Cuhls 2020                — horizon scanning
  Amanatidou et al. 2012    — horizon scanning
  Cresswell & Clark 2017    — multi-phase research design
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).parent.parent / "docs" / "Methods_Trends_2026.docx"

doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(3)


def heading(text, level=1):
    p = doc.add_heading(text, level=level)
    p.runs[0].font.size = Pt(12 if level == 1 else 11)
    return p


def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.runs[0]
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    return p


def italic_note(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.name = "Times New Roman"
    p.paragraph_format.space_after = Pt(4)
    return p


def add_search_table():
    """Insert a compact search parameter table for Strategy 1."""
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"

    headers = [("Parameter", "Detail")]
    rows_data = [
        ("Search date",    "10–12 August 2024"),
        ("Databases",      "OpenAlex (automated, n = 2,100); Google Scholar, "
                           "Elicit, Semantic Scholar (supplementary manual verification)"),
        ("Query strings",  '"Future Trend Impact Analysis"\n'
                           '"Future studies Fisheries Research"\n'
                           '"Future Fisheries Scenarios"\n'
                           '"Future of Fisheries"\n'
                           '"Future Trends and Drivers Fisheries"\n'
                           '"Megatrend Impacts Fisheries"\n'
                           '"Exploring Megatrends Fisheries"'),
        ("Results per query", "300 (OpenAlex API; polite pool, email provided)"),
    ]

    for i, (col_a, col_b) in enumerate(headers + rows_data):
        row = table.rows[i]
        row.cells[0].text = col_a
        row.cells[1].text = col_b
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
        if i == 0:
            for cell in row.cells:
                for run in cell.paragraphs[0].runs:
                    run.bold = True

    doc.add_paragraph()  # spacing after table


# ============================================================
# METHODS SECTION
# ============================================================

heading("3. Methods", level=1)

# ── 3.1 Methodological overview ──────────────────────────────
heading("3.1 Methodological overview", level=2)

body(
    "This review combined elements of a structured literature review (SLR) and horizon "
    "scanning to identify and analyse drivers, trends, and weak signals shaping the future "
    "of fisheries. The SLR component provided a rigorous and transparent framework for "
    "reducing bias in the search, selection, and synthesis of literature (Bearman et al., "
    "2012; Pati and Lorusso, 2017). The horizon scanning component guided the exploration "
    "of early indicators of change across the Social, Technological, Economic, Environmental, "
    "and Political (STEEP) domains, drawing on a constructivist notion of signal detection "
    "that treats weak signals as socially interpreted observations rather than objective "
    "facts about the future (Cuhls, 2020; Amanatidou et al., 2012). The integration of "
    "these two approaches is consistent with the multi-phase research design rationale "
    "described by Cresswell and Clark (2017), in which distinct methodological strategies "
    "address sequential research objectives within a single overarching inquiry."
)

body(
    "The review followed a three-step search strategy. Step 1 was a structured review of "
    "peer-reviewed academic literature addressing fisheries and futures studies. Step 2 was "
    "a manual search of cross-sectoral foresight publications to identify broader megatrends "
    "and drivers relevant to fisheries systems. Step 3 was an initial horizon scan for weak "
    "and emerging signals from grey literature and non-academic sources. Steps 1 and 2 were "
    "conducted simultaneously in August 2024; Step 3 constitutes an ongoing living document "
    "that is updated as new signals are identified. The overall search strategy followed "
    "PRISMA 2020 reporting guidelines (Page et al., 2021). A PRISMA flow diagram "
    "summarising search outcomes is provided in Figure 1."
)

# ── 3.2 Strategy 1 ───────────────────────────────────────────
heading("3.2 Search strategy 1: peer-reviewed fisheries futures literature", level=2)

body(
    "The purpose of the first search strategy was to compile and synthesise the existing "
    "peer-reviewed literature addressing future scenarios, impact analyses, and trend "
    "evaluations in fisheries research. The primary automated search was conducted via "
    "OpenAlex, a free and openly accessible academic database covering approximately "
    "250 million scholarly works across all disciplines (Priem et al., 2022 [VERIFY]). "
    "Seven query strings were applied (Table 1), each retrieving up to 300 results "
    "from the OpenAlex API using the polite pool protocol. Search terms were constructed "
    "to capture literature at the intersection of futures studies methodologies and "
    "fisheries research, and were iteratively tested for recall and precision prior to "
    "the search date."
)

italic_note("[Author note: Insert Table 1 — search parameters — here.]")
add_search_table()

body(
    "The 2,100 retrieved records were deduplicated using exact DOI matching followed "
    "by fuzzy title comparison (threshold: 90% similarity using the RapidFuzz library), "
    "which removed 479 duplicate records and yielded 1,621 unique entries. These were "
    "screened at title and abstract level by the primary reviewer against the following "
    "inclusion criteria: (1) the publication explicitly addresses future scenarios, drivers, "
    "or trend analyses in the context of fisheries or marine capture fisheries; and "
    "(2) the publication employs a future-oriented analytical perspective — including "
    "scenario analysis, projections, foresight methods, or trend assessment — rather than "
    "a purely retrospective or descriptive approach. Publications were excluded if they: "
    "(1) were not related to fisheries or marine capture fisheries; (2) lacked a "
    "future-oriented perspective (e.g. retrospective stock assessments); or (3) addressed "
    "only biological population dynamics without reference to broader social, economic, or "
    "governance drivers. Following title and abstract screening, 901 records were excluded "
    "and 720 were retained as potentially eligible; these 720 candidates are pending "
    "full-text eligibility assessment and are not included in the quantitative analyses "
    "reported in Section 4."
)

body(
    "In parallel with the automated search, 60 publications were identified and included "
    "through manual database searching and citation tracking conducted by the primary "
    "reviewer. These 60 manually identified papers constitute the fisheries futures "
    "evidence base for the analyses reported in Sections 4.2 and 4.3."
)

# ── 3.3 Strategy 2 ───────────────────────────────────────────
heading("3.3 Search strategy 2: cross-sectoral foresight publications", level=2)

body(
    "The second search strategy aimed to identify cross-sectoral drivers and megatrends "
    "from outside the fisheries-specific literature that may be consequential for fisheries "
    "systems over the medium to long term. Rather than a structured database search, this "
    "strategy employed a purposive manual scan of established megatrend foresight "
    "publications, conducted between 10 and 12 August 2024. Sources were selected on the "
    "basis of their institutional standing, citation prominence in the foresight literature, "
    "and coverage of global or multi-sectoral trends. The scan encompassed publications from "
    "major consultancy organisations (including PwC, KPMG, EY, and Deloitte), intergovernmental "
    "and governmental bodies (including the European Commission Joint Research Centre, the "
    "OECD, and national science and innovation agencies), and peer-reviewed academic journals "
    "addressing futures studies and megatrend analysis."
)

body(
    "For each source, individual trend or driver entries were extracted and recorded in the "
    "cross-sectoral megatrend database with standardised metadata: trend label, details, "
    "author, source, publication year, and source type. A total of 145 entries were retained "
    "across seven source type categories. Sources that addressed only a single narrow topic "
    "without a futures orientation, or that could not be reliably attributed to an identifiable "
    "institutional author, were excluded. No formal inter-rater reliability procedure was "
    "applied to this strategy; selection decisions were made by the primary reviewer and "
    "documented in the database comments field."
)

# ── 3.4 Strategy 3 ───────────────────────────────────────────
heading("3.4 Search strategy 3: weak and emerging signals", level=2)

body(
    "The third search strategy focused on identifying weak and emerging signals of change "
    "not yet prominent in formal academic or institutional literature. Following Hiltunen "
    "(2010) and Rossel (2021), weak signals are understood here as early, often ambiguous "
    "observations that, when interpreted through an appropriate analytical lens, may "
    "indicate incipient shifts in the conditions facing fisheries systems. This review "
    "adopts a constructivist position: signals are not objective facts about the future "
    "but are socially constructed sense-making hypotheses about nascent developments "
    "(Cuhls, 2020). Accordingly, signal identification relies on informed interpretive "
    "judgement rather than systematic database searching."
)

body(
    "Signals were identified through a combination of targeted grey literature review "
    "and active monitoring of technology, policy, and industry news sources relevant to "
    "fisheries and adjacent sectors. Each signal was documented in the living signal "
    "database with standardised metadata including: signal description, contextual details, "
    "source, publication date, and a reviewer commentary on potential fisheries relevance. "
    "The signal database is explicitly designed as an ongoing, collaborative resource "
    "rather than a completed systematic search; entries reported in this review should "
    "be understood as illustrative of the types of early-stage changes warranting "
    "monitoring, rather than as a comprehensive inventory."
)

# ── 3.5 Thematic classification ──────────────────────────────
heading("3.5 Thematic classification", level=2)

body(
    "All entries in the megatrend and fisheries futures databases were thematically "
    "classified using the STEEP framework (Social, Technological, Economic, Environmental, "
    "Political). Classification was performed computationally using a keyword-matching "
    "algorithm implemented in Python (source code available at: "
    "github.com/tmrmn/TrendFish2026). The algorithm concatenated the available text "
    "fields for each entry (trend label, details, and paper title), converted the "
    "combined string to lowercase, and applied a first-match rule against pre-specified "
    "keyword lists for each STEEP category. Entries that matched no keyword in any "
    "category, or for which all text fields were empty, were assigned to a residual "
    "Other/Unclassified category. The keyword lists were refined iteratively to improve "
    "precision, including the addition of fisheries-specific terminology "
    "(e.g. 'fishery', 'fisher', 'aquaculture', 'trophic') not present in generic STEEP "
    "schemes, and the replacement of overly generic terms (e.g. 'nature', which matched "
    "'changing nature of work') with more specific compound expressions "
    "(e.g. 'nature conservation', 'nature-based', 'natural resource'). Classification "
    "was not subject to formal inter-rater reliability testing; this is acknowledged "
    "as a methodological limitation in Section 5c."
)

body(
    "The gap analysis reported in Section 4.2 compares the STEEP distribution of the "
    "cross-sectoral megatrend database with that of the fisheries futures literature, "
    "expressed as percentage-point differences. Percentages were calculated over all "
    "entries in each database, including the Other/Unclassified category in the "
    "denominator. Entries in the Other/Unclassified category were excluded from the "
    "profile chart (Figure 5) to enable a cleaner visual comparison of the five "
    "classifiable STEEP domains. All analyses were conducted in Python using the "
    "pandas and matplotlib libraries; code and data are openly available in the "
    "project repository."
)

italic_note(
    "[Author note: Add the OpenAlex citation — Priem, J., Piwowar, H., and Orr, R. "
    "(2022). OpenAlex: A fully-open index of the world's research literature. "
    "arXiv:2205.01833. Verify before submission.]"
)

# ── Save ─────────────────────────────────────────────────────
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"Saved: {OUT}")
