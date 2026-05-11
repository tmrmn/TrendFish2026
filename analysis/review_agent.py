"""
review_agent.py — Independent scientist review of the manuscript draft.

Reads all section documents, sends the full text to Claude as a senior
peer-reviewer for Reviews in Fish Biology and Fisheries, and writes the
structured review to:
  docs/Review_Trends_2026.md    (plain text, easy to read)
  docs/Review_Trends_2026.docx  (formatted Word document)

Usage:
  python analysis/review_agent.py

API key:
  Set ANTHROPIC_API_KEY as an environment variable, or create a file
  called .env in the project root containing:
      ANTHROPIC_API_KEY=sk-ant-...
"""

import os
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
import anthropic

BASE = Path(__file__).parent.parent
DOCS = BASE / "docs"

# Section files in manuscript order (same as merge_manuscript.py)
SOURCES = [
    ("Abstract",      DOCS / "Abstract_Trends_2026.docx"),
    ("Introduction",  DOCS / "Introduction_Trends_2026.docx"),
    ("Methods",       DOCS / "Methods_Trends_2026.docx"),
    ("Results",       DOCS / "Results_Trends_2026.docx"),
    ("Discussion",    DOCS / "Discussion_Trends_2026.docx"),
    ("Conclusions",   DOCS / "Conclusions_Trends_2026.docx"),
]

OUT_MD   = DOCS / "Review_Trends_2026.md"
OUT_DOCX = DOCS / "Review_Trends_2026.docx"

JOURNAL  = "Reviews in Fish Biology and Fisheries"
MODEL    = "claude-opus-4-7"          # most capable model for deep review


# ── API key ───────────────────────────────────────────────────

def load_api_key() -> str:
    key = os.environ.get("ANTHROPIC_API_KEY", "")
    if not key:
        env_file = BASE / ".env"
        if env_file.exists():
            for line in env_file.read_text(encoding="utf-8").splitlines():
                if line.startswith("ANTHROPIC_API_KEY="):
                    key = line.split("=", 1)[1].strip().strip('"').strip("'")
                    break
    if not key:
        print(
            "ERROR: ANTHROPIC_API_KEY not found.\n"
            "  Option 1: set it as an environment variable.\n"
            "  Option 2: create a .env file in the project root:\n"
            "            ANTHROPIC_API_KEY=sk-ant-...\n"
        )
        sys.exit(1)
    return key


# ── Extract manuscript text ───────────────────────────────────

def extract_docx_text(path: Path) -> str:
    if not path.exists():
        return f"[SECTION MISSING: {path.name}]"
    doc = Document(path)
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    return "\n\n".join(lines)


def build_manuscript_text() -> str:
    parts = []
    for label, path in SOURCES:
        text = extract_docx_text(path)
        parts.append(f"## {label}\n\n{text}")
    return "\n\n---\n\n".join(parts)


# ── Reviewer prompt ───────────────────────────────────────────

SYSTEM_PROMPT = f"""You are a senior scientist acting as an independent peer reviewer for the journal \
*{JOURNAL}*. You have broad expertise in fisheries science, marine ecology, science policy, \
and systematic review methodology. You are thorough, constructive, and direct.

Your task is to review the manuscript draft provided and produce a structured review that \
mirrors what a journal editor would send to the author. Be honest about weaknesses — \
the author needs actionable feedback, not flattery.

Structure your review as follows:

1. **Summary** (3–5 sentences): What the paper does, its core argument, and its main contribution.

2. **Major Concerns** (numbered list): Issues that must be addressed before the paper can be \
accepted. These may concern conceptual framing, methodological rigour, missing analysis, \
unsupported claims, or logical gaps.

3. **Minor Concerns** (numbered list): Smaller issues — clarity, structure, citation gaps, \
language, consistency of terminology, formatting.

4. **Specific Comments** (by section): Go through each section (Abstract, Introduction, Methods, \
Results, Discussion, Conclusions) and provide targeted comments with line-level suggestions \
where relevant. Quote the original text when pointing to a specific passage.

5. **Recommendation**: One of — Accept, Minor Revision, Major Revision, Reject with \
encouragement to resubmit, Reject. Justify your recommendation in 2–3 sentences.

Be specific. Vague praise or vague criticism is unhelpful. If something is missing, say \
what it is and why it matters. If a section is strong, say so briefly and move on."""

USER_PROMPT_TEMPLATE = """Please review the following manuscript draft for *{journal}*.

---

{manuscript}

---

Provide your full peer review following the structure in your instructions."""


# ── Write Word output ─────────────────────────────────────────

def write_docx(review_text: str):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin    = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin   = Cm(3)
        sec.right_margin  = Cm(3)

    # Title
    p = doc.add_paragraph()
    r = p.add_run("Independent Peer Review")
    r.bold = True
    r.font.size = Pt(14)
    r.font.name = "Times New Roman"

    p = doc.add_paragraph()
    r = p.add_run(f"Manuscript: Drivers, Trends, and Signals in Fisheries Research\n"
                  f"Journal: {JOURNAL}\n"
                  f"Reviewer: AI independent review (claude-opus-4-7)")
    r.font.size = Pt(10)
    r.font.name = "Times New Roman"
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()

    # Body — parse markdown headings roughly into Word styles
    for line in review_text.splitlines():
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph()
            continue
        if stripped.startswith("## "):
            p = doc.add_paragraph(stripped[3:], style="Heading 2")
        elif stripped.startswith("# "):
            p = doc.add_paragraph(stripped[2:], style="Heading 1")
        elif stripped.startswith("**") and stripped.endswith("**") and len(stripped) < 80:
            p = doc.add_paragraph()
            r = p.add_run(stripped.strip("*"))
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
        else:
            p = doc.add_paragraph(stripped)
            p.style.font.name = "Times New Roman"
            p.style.font.size = Pt(11)

    doc.save(OUT_DOCX)
    print(f"Saved Word document: {OUT_DOCX}")


# ── Main ──────────────────────────────────────────────────────

def main():
    api_key = load_api_key()

    print("Reading manuscript sections…")
    manuscript = build_manuscript_text()
    word_count = len(manuscript.split())
    print(f"  Total words extracted: {word_count:,}")

    print(f"\nSending to {MODEL} for review (streaming)…\n")
    print("=" * 72)

    client = anthropic.Anthropic(api_key=api_key)

    review_chunks = []
    with client.messages.stream(
        model=MODEL,
        max_tokens=4096,
        system=SYSTEM_PROMPT,
        messages=[
            {
                "role": "user",
                "content": USER_PROMPT_TEMPLATE.format(
                    journal=JOURNAL,
                    manuscript=manuscript,
                ),
            }
        ],
    ) as stream:
        for text in stream.text_stream:
            print(text, end="", flush=True)
            review_chunks.append(text)

    print("\n" + "=" * 72)

    review_text = "".join(review_chunks)

    # Save markdown
    OUT_MD.write_text(
        f"# Peer Review — Drivers, Trends, and Signals in Fisheries Research\n\n"
        f"*Journal: {JOURNAL}*  \n"
        f"*Reviewer: AI independent review (claude-opus-4-7)*\n\n---\n\n"
        + review_text,
        encoding="utf-8",
    )
    print(f"\nSaved Markdown:      {OUT_MD}")

    # Save Word doc
    write_docx(review_text)


if __name__ == "__main__":
    main()
