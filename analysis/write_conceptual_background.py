"""
Generate Section 2 — Conceptual Background as a formatted Word document.
Writes to: ../docs/ConceptualBackground_Trends_2026.docx
Target: ~700 words.

Covers:
  2.1 Futures studies and horizon scanning
  2.2 Drivers, trends, megatrends, and weak signals
  2.3 The STEEP framework
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).parent.parent / "docs" / "ConceptualBackground_Trends_2026.docx"

doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(3)


def heading(text, level=1):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12 if level == 1 else 11)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after  = Pt(4)
    return p


def body(text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in p.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
    p.paragraph_format.space_after  = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0.75)
    return p


# ── Section heading ───────────────────────────────────────────
heading("2. Conceptual Background", level=1)

# ── 2.1 ──────────────────────────────────────────────────────
heading("2.1 Futures studies and horizon scanning", level=2)

body(
    "Futures studies is an interdisciplinary field concerned with the systematic exploration "
    "of possible, probable, and preferable futures (Bell, 2003). It encompasses a diverse "
    "repertoire of methods, including scenario analysis, Delphi surveys, trend extrapolation, "
    "causal layered analysis, and horizon scanning, each suited to different types of foresight "
    "question and different degrees of uncertainty. What unifies these approaches is a shared "
    "orientation towards the structured analysis of change over medium to long-term timescales, "
    "and a concern with making the implicit assumptions embedded in policy and planning decisions "
    "explicit and open to revision (Cuhls, 2020)."
)

body(
    "Horizon scanning is a specific form of futures inquiry focused on the systematic detection "
    "of early indicators of change — including emerging trends, novel technologies, and weak "
    "signals — that lie at or beyond the current frontier of public attention (Sutherland et al., "
    "2011). Unlike scenario planning, which constructs coherent narratives about future states, "
    "horizon scanning is primarily concerned with monitoring the periphery of current knowledge "
    "for developments that may become significant over time. It has been applied across a range "
    "of policy domains — including public health (Fisher et al., 2016), biodiversity conservation "
    "(Sutherland and Woodroof, 2009), and food systems (Tàbara et al., 2017) — but remains "
    "comparatively underdeveloped in fisheries research."
)

# ── 2.2 ──────────────────────────────────────────────────────
heading("2.2 Drivers, trends, megatrends, and weak signals", level=2)

body(
    "This review employs four conceptual categories to organise the evidence base, each "
    "representing a distinct scale and degree of certainty in the analysis of change. These "
    "categories are positioned along a spectrum from broad, well-established forces operating "
    "at civilisational scale, to early and uncertain signals at the margins of formal knowledge."
)

body(
    "Drivers of change are the underlying forces or pressures that shift conditions within a "
    "system, often operating at scales beyond the immediate control of individual stakeholders "
    "but with direct consequences for how institutions, communities, and industries function "
    "(Vecchiato and Roveda, 2010). Drivers may be biophysical — such as ocean warming or "
    "nutrient loading — or socially constituted, such as regulatory change, market restructuring, "
    "or shifts in consumer preferences. They provide the causal substrate from which observable "
    "trends emerge."
)

body(
    "Trends are widespread, directional changes that unfold over time and influence attitudes, "
    "policies, and economic activity at societal scale (Celi and Colombi, 2020). They represent "
    "the observable manifestation of underlying drivers, and are typically sustained over "
    "multi-year periods before becoming institutionally recognised. When trends operate at a "
    "global scale and over multi-decadal timescales — reshaping fundamental conditions across "
    "multiple sectors simultaneously — they are commonly described as megatrends: powerful "
    "forces expected to bring about significant transformations in societal or natural "
    "environments over a period of five to twenty or more years (Naisbitt, 1982; Hajkowicz, "
    "2015). Megatrends are distinguished from ordinary trends by their breadth, duration, and "
    "the degree to which they cut across sectoral, disciplinary, and geographical boundaries."
)

body(
    "Weak signals occupy the opposite end of the certainty spectrum. Following Hiltunen (2010) "
    "and Rossel (2021), weak signals are understood here as early, often ambiguous observations "
    "that, when interpreted through an appropriate analytical lens, may indicate incipient shifts "
    "in the conditions facing fisheries systems. They do not yet appear in formal datasets or "
    "policy agendas, and their eventual significance is uncertain; their value lies precisely in "
    "the advance notice they may provide to practitioners willing to engage with uncertainty "
    "constructively. This review adopts a constructivist position with respect to signal "
    "detection: signals are not objective facts about the future, but socially constructed "
    "sense-making hypotheses about nascent developments (Cuhls, 2020)."
)

# ── 2.3 ──────────────────────────────────────────────────────
heading("2.3 The STEEP framework", level=2)

body(
    "The STEEP framework — an acronym for Social, Technological, Economic, Environmental, and "
    "Political — is a widely used organising scheme in strategic foresight and environmental "
    "scanning (Aguilar, 1967; Morrison, 1992). It provides a structured vocabulary for "
    "categorising the domains from which drivers and trends emerge, and facilitates systematic "
    "comparison across different evidence bases. The framework has been applied in megatrend "
    "analyses across multiple sectors (Kalaitzi et al., 2020; Malik and Janowska, 2021) and "
    "in policy foresight processes at the European Commission (Knowledge4Policy, 2022) and "
    "national science agencies (Naughtin et al., 2022)."
)

body(
    "In this review, STEEP classification serves two analytical functions. First, it provides "
    "a common taxonomic basis for comparing the thematic coverage of the cross-sectoral "
    "megatrend literature with that of the fisheries-specific futures literature — enabling the "
    "identification of disciplinary gaps where fisheries research has systematically "
    "under-engaged with particular domains of change. Second, it structures the presentation "
    "of results in Section 4, organising drivers and trends by the domain from which they "
    "primarily originate. Classification was applied computationally using a keyword-matching "
    "algorithm (see Section 3.5), with entries that did not match any STEEP keyword retained "
    "in a residual Other/Unclassified category."
)

# ── Save ──────────────────────────────────────────────────────
doc.save(OUT)
print(f"Saved: {OUT}")
wc = sum(len(p.text.split()) for p in doc.paragraphs if p.text.strip())
print(f"Word count (approx): {wc}")
