"""
Generate Discussion section as a formatted Word document.
Writes to: ../docs/Discussion_Trends_2026.docx
Target: ~1,200 words across 5a (interpretation of results) and 5b (implications).
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm

OUT = Path(__file__).parent.parent / "docs" / "Discussion_Trends_2026.docx"

doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(3)


def heading(text, level=1):
    p = doc.add_heading(text, level=level)
    p.runs[0].font.size = Pt(12 if level == 1 else 11)
    return p


def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.runs[0]
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    return p


def italic_note(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.name = "Times New Roman"
    p.paragraph_format.space_after = Pt(4)
    return p


# ============================================================
# DISCUSSION SECTION
# ============================================================

heading("5. Discussion", level=1)

# ── 5a Discussion of Results ──────────────────────────────────
heading("5a. Interpretation of findings", level=2)

body(
    "The results of this review reveal a systematic and consequential asymmetry between "
    "the futures landscape as characterised in the cross-sectoral megatrend literature "
    "and the futures landscape as reflected in fisheries-specific research. Whilst both "
    "bodies of literature converge on climate change as the dominant driver of change — "
    "a finding that is consistent with the broader futures studies canon (Hajkowicz, 2015; "
    "Naughtin et al., 2022) — the fisheries futures literature exhibits a markedly narrower "
    "disciplinary range. The concentration of fisheries-specific entries in the environmental "
    "domain (43%) is, on one level, unsurprising: ecological disruption, overfishing, "
    "by-catch, and habitat loss are longstanding concerns in fisheries science, and the "
    "literature accordingly reflects sustained engagement with these pressures. Yet this "
    "concentration simultaneously reveals an absence."
)

body(
    "The most striking finding of this review is the complete absence of technological "
    "drivers within the fisheries futures literature. Across the cross-sectoral megatrend "
    "database, technological change — encompassing accelerating automation, hyperconnectivity, "
    "synthetic biology, and digitalisation — constituted 28% of all classified entries, "
    "representing one of the two largest driver categories alongside environmental change. "
    "No comparable body of work was identified in the fisheries-specific literature "
    "(0%, gap = +28 percentage points). This gap exceeds the analogous deficits observed "
    "in the social and economic domains — themselves substantial at +11 and +14 percentage "
    "points respectively — and represents the largest disciplinary blind spot identified "
    "in this review. The finding suggests that the structural implications of technological "
    "transformation for fisheries systems — including the automation of fleet operations, "
    "the rise of remote sensing and machine learning in stock assessment, the emergence of "
    "alternative protein sources competing commercially with wild-catch seafood, and shifting "
    "labour dynamics within fishing communities — have not yet been systematically integrated "
    "into futures-oriented fisheries research."
)

body(
    "The underrepresentation of social and economic drivers in the fisheries futures "
    "literature is likewise notable, though the mechanisms may differ. Social drivers — "
    "including demographic change, urbanisation, and shifting consumer preferences — "
    "are well-established in the cross-sectoral foresight literature as significant forces "
    "reshaping demand for seafood and the composition of fishing communities (KPMG, 2014; "
    "Knowledge4Policy, 2022). Their limited presence in fisheries-specific futures research "
    "may reflect disciplinary conventions that situate social dynamics as context for "
    "ecological or management analyses, rather than as primary objects of foresight inquiry. "
    "Economic foresight — addressing market restructuring, the growing economic weight of "
    "emerging markets, and the globalisation of seafood supply chains — was similarly "
    "underrepresented, despite the centrality of economic incentives to fisheries governance "
    "and sustainability transitions."
)

body(
    "A further observation concerns the temporal and methodological range of the fisheries "
    "futures literature. The majority of identified studies employed scenario analysis or "
    "quantitative projection techniques to model ecological outcomes under specified "
    "assumptions. Whilst this approach has produced important insights — particularly in "
    "relation to the projected effects of climate change on fish stock distributions and "
    "yields (Cheung et al., 2010; Lotze et al., 2019) — it represents only a subset of the "
    "foresight methods available. Approaches such as horizon scanning, Delphi methods, "
    "causal layered analysis, and participatory futures workshops, which are routinely "
    "employed in cross-sectoral foresight practice, appear sparingly in the fisheries "
    "literature. This methodological narrowness may itself constrain the range of futures "
    "that the field is able to articulate and act upon."
)

# ── 5b Discussion of Implications ────────────────────────────
heading("5b. Implications for research, education, and policy", level=2)

body(
    "For researchers, the findings point to a clear priority: the development of "
    "futures-oriented fisheries research that takes technological change seriously as a "
    "driver in its own right. This does not require the wholesale adoption of technological "
    "determinism, but rather the integration of technology assessment and foresight into the "
    "analytical repertoire of fisheries science. Particular attention is warranted for "
    "emerging technologies that are already transforming adjacent sectors but whose "
    "implications for fisheries remain poorly characterised: these include precision "
    "fermentation and cultivated seafood, advanced monitoring systems for illegal, "
    "unreported, and unregulated (IUU) fishing, and AI-driven stock assessment. More "
    "broadly, the review underscores the value of interdisciplinary collaboration between "
    "fisheries scientists, foresight practitioners, economists, and social scientists. "
    "Futures studies is an inherently transdisciplinary field (Cuhls, 2020), and its "
    "application to fisheries is likely to yield richer results through such collaboration "
    "than through disciplinary extension alone. The living database introduced in this "
    "review — combining structured literature evidence with an ongoing signal-scanning "
    "protocol — is intended as a practical contribution to this integrative effort, offering "
    "a shared resource that the research community can extend and refine over time."
)

body(
    "For educators, particularly those involved in fisheries science training and marine "
    "resource management curricula, the review highlights a curricular gap. The relative "
    "absence of social, economic, and technological foresight in the fisheries futures "
    "literature suggests that students and early-career researchers in the sector may have "
    "limited exposure to the conceptual tools of futures studies — including the distinction "
    "between drivers, trends, and signals; the use of megatrend frameworks such as STEEP; "
    "and the methodological diversity of horizon scanning and scenario planning. Incorporating "
    "these approaches into graduate-level training in fisheries science would better equip "
    "future researchers and practitioners to anticipate, rather than merely respond to, "
    "structural change in the sector. The conceptual vocabulary introduced in this review "
    "is intended to serve, in part, as a resource for such curricular development."
)

body(
    "For policymakers and managers, the most immediate implication of this review concerns "
    "the risk of an evidence base that is systematically skewed towards environmental "
    "drivers. Fisheries management decisions are increasingly taken in contexts shaped by "
    "forces — technological disruption, geopolitical realignment, demographic change, and "
    "economic restructuring — that lie at the margins of the existing fisheries futures "
    "literature. Governance frameworks designed principally to address stock depletion and "
    "ecological sustainability may be poorly positioned to respond to these emerging "
    "pressures, including the governance of automated fishing fleets, the regulation of "
    "novel seafood substitutes, or the distributional consequences of shifting seafood "
    "supply chains for small-scale fishing communities. Bridging the gap between "
    "cross-sectoral foresight intelligence and fisheries policy requires sustained "
    "institutional investment in horizon scanning capacity within fisheries management "
    "bodies — a function for which the database compiled in this review provides an "
    "initial framework."
)

# ── Limitations ───────────────────────────────────────────────
heading("5c. Limitations", level=2)

body(
    "Several limitations of this review warrant acknowledgement. The cross-sectoral "
    "megatrend database was assembled through a manual search of known foresight "
    "publications and grey literature sources, and cannot be considered exhaustive; "
    "the selection of sources necessarily reflects the availability of English-language "
    "institutional reports and the reviewers' initial knowledge of the foresight landscape. "
    "The STEEP classification scheme, whilst well-established in futures studies, involves "
    "a degree of interpretive judgement: many drivers are inherently multidimensional, and "
    "classification by dominant characteristic inevitably simplifies cross-cutting phenomena. "
    "The weak signal database remains in its initial stages, with systematic signal scanning "
    "yet to be completed; the signals identified at this stage are indicative rather than "
    "representative. Finally, the automated OpenAlex search yielded 720 candidate papers "
    "pending full-text eligibility assessment: the final composition of the fisheries "
    "futures evidence base may be revised as this assessment proceeds."
)

italic_note(
    "[Author note: Consider adding a sentence on the language bias (English-only) "
    "and whether non-English foresight literature — particularly from East and Southeast "
    "Asia, where fisheries futures research is active — may affect the gap analysis findings.]"
)

# ── Save ──────────────────────────────────────────────────────
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"Saved: {OUT}")
